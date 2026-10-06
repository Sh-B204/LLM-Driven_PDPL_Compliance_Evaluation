"""Synthetic demo data (no real policies): used by --mock runs and the test-suite."""
from __future__ import annotations

from pathlib import Path

DEMO_PDPL = [
    ("Article 4 - Rights of the Data Subject", "The data subject has the right to be informed, to access, to request correction and to request destruction of personal data."),
    ("Article 5 - Consent", "Personal data may not be processed without the consent of the data subject, who may withdraw consent at any time."),
    ("Article 10 - Collection", "Collection shall be limited to the minimum data necessary for the stated purpose."),
    ("Article 19 - Security", "The controller shall take organisational, administrative and technical measures to protect personal data."),
    ("Article 20 - Breach notification", "The controller shall notify the competent authority when a leak of personal data occurs."),
    ("Article 29 - Transfers", "Personal data may not be transferred outside the Kingdom except under the conditions of the regulation."),
]
SHORT = [
    ("Introduction to this policy", "This privacy policy explains how we collect and use your personal data."),
    ("Your rights", "You may access, correct or delete your data by contacting us. You can object to processing."),
    ("\u0627\u0644\u0623\u0645\u0627\u0646", "\u0646\u062d\u0646 \u0646\u062a\u062e\u0630 \u0625\u062c\u0631\u0627\u0621\u0627\u062a \u062a\u0642\u0646\u064a\u0629 \u0644\u062d\u0645\u0627\u064a\u0629 \u0628\u064a\u0627\u0646\u0627\u062a\u0643 \u0627\u0644\u0634\u062e\u0635\u064a\u0629."),
    ("Retention", "We keep data only as long as necessary for the purposes described."),
]


def _write(path: Path, sections, intro=None):
    from docx import Document
    d = Document()
    if intro:
        d.add_paragraph(intro)
    for title, body in sections:
        d.add_heading(title, level=1)
        d.add_paragraph(body)
    path.parent.mkdir(parents=True, exist_ok=True)
    d.save(str(path))


def create_demo_data(target: Path) -> dict:
    target = Path(target)
    long_sections = [(f"Section {i}: topic {i}", (f"Paragraph about data practice {i}. " * 12).strip()) for i in range(1, 16)]
    _write(target / "DemoPDPL.docx", DEMO_PDPL)
    _write(target / "DemoShort.docx", SHORT)
    _write(target / "DemoLong.docx", long_sections)
    return {"pdpl": target / "DemoPDPL.docx", "short": target / "DemoShort.docx", "long": target / "DemoLong.docx"}
