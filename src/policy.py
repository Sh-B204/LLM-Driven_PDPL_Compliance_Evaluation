"""Policy / PDPL DOCX loading (notebook cells 8 and 10, logic unchanged)."""
from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Dict, List, Tuple

from .utils import sha256_file


def parse_docx_to_sections(docx_path: str | Path) -> List[Dict]:
    """Notebook parse_docx_to_sections: split by paragraphs whose style name contains 'Heading'."""
    from docx import Document
    doc = Document(str(docx_path))
    sections: List[Dict] = []
    current_section = {"title": "Introduction", "content": ""}
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        if "Heading" in para.style.name:
            if current_section["content"]:
                sections.append(current_section)
            current_section = {"title": text, "content": ""}
        else:
            current_section["content"] += " " + text
    if current_section["content"]:
        sections.append(current_section)
    return sections


def parse_docx_to_texts_metadatas(docx_path: str | Path) -> Tuple[List[str], List[Dict]]:
    """Notebook parse_docx_to_texts_metadatas: PDPL reference chunks (one per heading section)."""
    from docx import Document as Docx
    doc = Docx(str(docx_path))
    texts: List[str] = []
    metadatas: List[Dict] = []
    current_section_content = ""
    current_section_title = "Introduction"
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = para.style.name
        if "Heading" in style:
            if current_section_content:
                texts.append(current_section_content.strip())
                metadatas.append({"title": current_section_title, "content": str(docx_path)})
            current_section_title = text
            current_section_content = ""
        else:
            current_section_content += " " + text
    if current_section_content:
        texts.append(current_section_content.strip())
        metadatas.append({"title": current_section_title, "content": str(docx_path)})
    return texts, metadatas


def build_full_policy_text(sections: List[Dict]) -> str:
    """Notebook: "\\n".join(f"Section {i}: {title}\\n{content}" ...)."""
    return "\n".join(f"Section {i}: {s['title']}\n{s['content']}" for i, s in enumerate(sections, 1))


@dataclass
class PolicyData:
    path: Path
    sha256: str
    sections: List[Dict]
    full_text: str

    @property
    def n_sections(self) -> int:
        return len(self.sections)


def load_policy(path: str | Path) -> PolicyData:
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"Policy file not found: {path}")
    sections = parse_docx_to_sections(path)
    if not sections:
        raise ValueError(f"No sections parsed from policy (empty or no text): {path}")
    return PolicyData(path=path, sha256=sha256_file(path), sections=sections,
                      full_text=build_full_policy_text(sections))
